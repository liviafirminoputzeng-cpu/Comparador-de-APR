"""Área compartilhada de acompanhamento de BowTies, separada do comparador."""

import base64
import hashlib
import io
from datetime import date
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from acompanhamento_db import (
    EdicaoConcorrente, SemPermissao, atualizar_item, cadastrar_itens, criar_projeto,
    email_normalizado, listar_historico, listar_itens, listar_membros,
    listar_projetos, papel_no_projeto, salvar_membro, SITUACOES,
)
from comparacao import candidato_bowtie
from extracao import extrair_cenarios_pdf


def _erro(erro):
    if isinstance(erro, EdicaoConcorrente):
        st.warning(str(erro))
    elif isinstance(erro, (ValueError, SemPermissao)):
        st.error(str(erro))
    else:
        st.error("Não foi possível acessar o acompanhamento. Verifique a conexão e o esquema do banco.")


def _ler_controle_excel(conteudo):
    """Lê os campos de identificação da planilha modelo, sem supor atribuições."""
    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    try:
        if "Controle" not in wb.sheetnames:
            raise ValueError("A planilha precisa ter a aba Controle do modelo de acompanhamento.")
        ws = wb["Controle"]
        esperados = ["ID do cenário", "Nó", "Cenário", "Sistema", "Evento topo / desvio"]
        encontrados = [ws.cell(8, col).value for col in range(1, 6)]
        if encontrados != esperados:
            raise ValueError("Os cabeçalhos da aba Controle não correspondem ao modelo.")
        itens = []
        for linha in ws.iter_rows(min_row=9, max_col=18, values_only=True):
            if not linha[0]:
                continue
            if any(linha[i] not in (None, "", "A definir") for i in range(5, 17)):
                raise ValueError(
                    "Esta planilha já contém atribuições ou andamento. A importação "
                    "automática de trabalho em curso exige conferência; nenhum dado foi importado."
                )
            itens.append({
                "ID": linha[0], "Nó": linha[1], "Cenário": linha[2],
                "Sistema": linha[3], "Perigo": linha[4], "Página PDF": linha[17],
            })
        if not itens:
            raise ValueError("Nenhum BowTie foi encontrado na aba Controle.")
        if len(itens) > 2000:
            raise ValueError("A planilha contém mais de 2.000 itens.")
        return itens
    finally:
        wb.close()


def _painel(url, projeto_id, email, administradores, chave):
    @st.fragment(run_every="20s")
    def conteudo():
        try:
            registros = listar_itens(url, projeto_id, email, administradores)
        except Exception as erro:
            _erro(erro)
            return
        st.caption("Painel sincronizado automaticamente a cada 20 segundos enquanto esta página estiver aberta.")
        total = len(registros)
        feitos = sum(r["situacao"] == "Concluído" for r in registros)
        em_andamento = sum(r["situacao"] in ("Em desenvolvimento", "Em revisão", "Correção solicitada") for r in registros)
        a_fazer = sum(r["situacao"] in ("A definir", "A fazer") for r in registros)
        bloqueados = sum(r["situacao"] == "Bloqueado" for r in registros)
        atrasados = sum(r["prazo"] is not None and r["prazo"] < date.today()
                         and r["situacao"] != "Concluído" for r in registros)
        cols = st.columns(6)
        for col, rotulo, valor in zip(cols,
            ("Total", "A fazer / definir", "Em andamento", "Concluídos", "Bloqueados", "Atrasados"),
            (total, a_fazer, em_andamento, feitos, bloqueados, atrasados)):
            col.metric(rotulo, valor)
        por_responsavel = {}
        for r in registros:
            chave_pessoa = r["responsavel_email"] or "Não atribuído"
            por_responsavel.setdefault(chave_pessoa, {"Total": 0, "Concluídos": 0, "Em andamento": 0})
            por_responsavel[chave_pessoa]["Total"] += 1
            por_responsavel[chave_pessoa]["Concluídos"] += (r["situacao"] == "Concluído")
            por_responsavel[chave_pessoa]["Em andamento"] += r["situacao"] in (
                "Em desenvolvimento", "Em revisão", "Correção solicitada"
            )
        if por_responsavel:
            st.subheader("Distribuição da equipe")
            st.dataframe(pd.DataFrame.from_dict(por_responsavel, orient="index")
                         .rename_axis("Responsável").reset_index(),
                         hide_index=True, use_container_width=True)
        st.subheader("BowTies do projeto")
        filtro_status = st.multiselect("Situação", list(SITUACOES), key=f"situacao_{chave}")
        filtro_no = st.multiselect(
            "Nó", sorted({str(r["no"]) for r in registros}), key=f"no_{chave}"
        )
        filtrados = [r for r in registros if
                     (not filtro_status or r["situacao"] in filtro_status) and
                     (not filtro_no or str(r["no"]) in filtro_no)]
        campos = {
            "id_cenario": "ID do cenário", "no": "Nó", "cenario": "Cenário",
            "sistema": "Sistema", "evento_topo": "Evento topo / desvio",
            "trabalho": "Trabalho", "responsavel_email": "Responsável",
            "situacao": "Situação", "prioridade": "Prioridade", "prazo": "Prazo",
            "progresso": "Progresso (%)", "proxima_acao": "Próxima ação / impedimento",
            "atualizado_em": "Atualizado em",
        }
        st.dataframe(pd.DataFrame([{titulo: r[campo] for campo, titulo in campos.items()}
                                   for r in filtrados], columns=list(campos.values())),
                     hide_index=True, use_container_width=True, height=440)

    conteudo()


def _editar(url, projeto_id, email, administradores, coordenador, membros, itens):
    st.subheader("Atualizar um BowTie")
    editaveis = itens if coordenador else [i for i in itens if i["responsavel_email"] == email]
    if not editaveis:
        st.info("Não há BowTies atribuídos a você neste projeto." if not coordenador
                else "Cadastre os BowTies para começar a distribuição.")
        return
    selecionado = st.selectbox(
        "Selecione pelo ID, nó e evento topo", editaveis,
        format_func=lambda i: f'{i["id_cenario"]} · Nó {i["no"]} · {i["evento_topo"]}',
        key=f"editar_{projeto_id}",
    )
    analistas = [m for m in membros if m["papel"] == "Analista"]
    emails = [m["email"] for m in analistas]
    # A versão de registro exibida no formulário é verificada no banco ao salvar.
    with st.form(f'form_item_{selecionado["id"]}_{selecionado["versao"]}'):
        st.caption(f'Fonte: {selecionado["fonte"] or "não informada"} · '
                   f'Página PDF: {selecionado["pagina_pdf"] or "não informada"}')
        if coordenador:
            trabalho = st.selectbox("Trabalho", [None, "Novo", "Correção"],
                                    index=[None, "Novo", "Correção"].index(selecionado["trabalho"]),
                                    format_func=lambda x: x or "A definir")
            opcoes = [None] + emails
            atual = selecionado["responsavel_email"]
            responsavel = st.selectbox("Responsável", opcoes, index=opcoes.index(atual)
                                      if atual in opcoes else 0,
                                      format_func=lambda x: x or "Não atribuído")
            prioridade = st.selectbox("Prioridade", [None, "Alta", "Média", "Baixa"],
                                     index=[None, "Alta", "Média", "Baixa"].index(selecionado["prioridade"]),
                                     format_func=lambda x: x or "A definir")
        situacoes_permitidas = list(SITUACOES) if coordenador else [s for s in SITUACOES if s != "Concluído"]
        situacao = st.selectbox("Situação", situacoes_permitidas,
                               index=situacoes_permitidas.index(selecionado["situacao"])
                               if selecionado["situacao"] in situacoes_permitidas else 0)
        progresso = st.number_input("Progresso (%)", min_value=0, max_value=100,
                                   value=selecionado["progresso"], placeholder="Ainda não informado")
        proxima_acao = st.text_area("Próxima ação / impedimento",
                                   value=selecionado["proxima_acao"] or "")
        if coordenador:
            inicio_previsto = st.date_input("Início previsto", value=selecionado["inicio_previsto"])
            prazo = st.date_input("Prazo", value=selecionado["prazo"])
        inicio_real = st.date_input("Início real", value=selecionado["inicio_real"])
        entrega_inicial = st.date_input("Entrega inicial", value=selecionado["entrega_inicial"])
        if coordenador:
            conclusao_validada = st.date_input("Conclusão validada", value=selecionado["conclusao_validada"])
        if st.form_submit_button("Salvar atualização", type="primary"):
            dados = {
                "situacao": situacao, "progresso": progresso,
                "proxima_acao": proxima_acao.strip(), "inicio_real": inicio_real,
                "entrega_inicial": entrega_inicial,
            }
            if coordenador:
                dados.update({
                    "trabalho": trabalho, "responsavel_email": responsavel,
                    "prioridade": prioridade, "inicio_previsto": inicio_previsto,
                    "prazo": prazo, "conclusao_validada": conclusao_validada,
                })
            try:
                alterado = atualizar_item(url, projeto_id, selecionado["id"], email,
                                         administradores, selecionado["versao"], dados)
                if alterado:
                    st.toast("Atualização salva para a equipe.")
                    st.rerun()
                else:
                    st.info("Nenhuma informação foi alterada.")
            except Exception as erro:
                _erro(erro)


def _cadastro(url, projeto_id, email, administradores):
    st.subheader("Cadastrar BowTies neste projeto")
    st.caption("Os dados técnicos vêm apenas da APR ou da planilha modelo. "
               "Uma importação repetida preserva os registros já existentes.")
    tab_pdf, tab_excel, tab_manual = st.tabs(["APR atualizada (PDF)", "Planilha modelo (Excel)", "Cadastro manual"])
    with tab_pdf:
        pdf = st.file_uploader("Envie a APR atualizada", type="pdf", key=f"apr_tracker_{projeto_id}")
        if pdf and st.button("Ler candidatos a BowTie", key=f"ler_pdf_{projeto_id}"):
            try:
                resultado = extrair_cenarios_pdf(pdf.getvalue())
                candidatos = [c for c in resultado["cenarios"] if candidato_bowtie(c)]
                st.session_state[f"candidatos_{projeto_id}"] = (
                    hashlib.sha256(pdf.getvalue()).hexdigest(), candidatos, pdf.name
                )
                st.success(f"{len(candidatos)} candidatos encontrados. Confira antes de cadastrar.")
            except Exception as erro:
                _erro(erro)
        dados_pdf = st.session_state.get(f"candidatos_{projeto_id}")
        if pdf and dados_pdf and dados_pdf[0] == hashlib.sha256(pdf.getvalue()).hexdigest():
            _, candidatos, fonte = dados_pdf
            st.dataframe(pd.DataFrame([{
                "ID": c.get("ID"), "Nó": c.get("Nó"), "Cenário": c.get("Cenário"),
                "Sistema": c.get("Sistema"), "Evento topo": c.get("Perigo"),
            } for c in candidatos]), use_container_width=True, hide_index=True)
            if candidatos and st.button("Cadastrar estes candidatos no projeto", key=f"import_pdf_{projeto_id}"):
                try:
                    novos, existentes = cadastrar_itens(url, projeto_id, email, administradores,
                                                        candidatos, fonte)
                    st.toast(f"{novos} cadastrados; {existentes} já existiam.")
                    st.rerun()
                except Exception as erro:
                    _erro(erro)
    with tab_excel:
        excel = st.file_uploader("Envie a planilha de acompanhamento modelo", type="xlsx",
                                 key=f"excel_tracker_{projeto_id}")
        if excel:
            try:
                candidatos = _ler_controle_excel(excel.getvalue())
                st.info(f"{len(candidatos)} itens identificados. Esta importação serve para "
                        "a planilha inicial, ainda sem responsáveis ou andamento preenchidos.")
                if st.button("Cadastrar itens da planilha", key=f"import_excel_{projeto_id}"):
                    novos, existentes = cadastrar_itens(url, projeto_id, email, administradores,
                                                        candidatos, excel.name)
                    st.toast(f"{novos} cadastrados; {existentes} já existiam.")
                    st.rerun()
            except Exception as erro:
                _erro(erro)
    with tab_manual:
        with st.form(f"cadastro_manual_{projeto_id}"):
            id_cenario = st.text_input("ID do cenário")
            no = st.text_input("Nó")
            cenario = st.text_input("Cenário")
            sistema = st.text_input("Sistema")
            perigo = st.text_area("Evento topo / desvio")
            fonte = st.text_input("Fonte do registro")
            if st.form_submit_button("Cadastrar BowTie"):
                try:
                    novos, existentes = cadastrar_itens(url, projeto_id, email, administradores,
                        [{"ID": id_cenario, "Nó": no, "Cenário": cenario,
                          "Sistema": sistema, "Perigo": perigo}], fonte)
                    if existentes:
                        st.warning("Este ID de cenário já existe neste projeto.")
                    else:
                        st.toast("BowTie cadastrado.")
                        st.rerun()
                except Exception as erro:
                    _erro(erro)


def mostrar_acompanhamento():
    imagem = Path(__file__).resolve().parent / "assets" / "putz_identidade.webp"
    capa = ("data:image/webp;base64," + base64.b64encode(imagem.read_bytes()).decode("ascii")) \
        if imagem.exists() else ""
    st.markdown("""<style>
    .stApp {background: #f3f5fc; color: #101632;}
    .block-container {max-width: 1500px; padding-top: 1.4rem;}
    .tracking-header {padding: 1.5rem 2rem; border-radius: 16px;
      background: linear-gradient(90deg,rgba(7,8,30,.95),rgba(7,8,30,.75),rgba(7,8,30,.2)),
      url('""" + capa + """') center / cover #101936;
      color: white; margin: .8rem 0 1.5rem;}
    .tracking-header h1 {color: white; margin: .2rem 0;}
    .tracking-header p {color: #dfecff; margin: .2rem 0;}
    </style><div class="tracking-header"><h1>Acompanhamento de BowTies</h1>
    <p>Projetos, distribuição de tarefas e andamento compartilhado da equipe.</p></div>""",
                unsafe_allow_html=True)
    try:
        configuracao = dict(st.secrets)
    except FileNotFoundError:
        configuracao = {}
    if "tracker_database_url" not in configuracao or "auth" not in configuracao:
        st.warning("O acompanhamento compartilhado ainda precisa de banco de dados e login "
                   "configurados nos Secrets do Streamlit. Consulte o GUIA_ACOMPANHAMENTO.md. "
                   "O comparador de APR continua disponível.")
        return
    if not st.user.is_logged_in:
        st.button("Entrar com a conta da equipe", on_click=st.login, type="primary")
        return
    email = email_normalizado(st.user.get("email"))
    if not email:
        st.error("O provedor de login não forneceu um e-mail. Confira a configuração de autenticação.")
        return
    administradores = {email_normalizado(x) for x in configuracao.get("tracker_admin_emails", [])}
    url = configuracao["tracker_database_url"]
    st.caption(f"Conectado como {email}")
    if st.button("Sair do acompanhamento", on_click=st.logout):
        return
    try:
        projetos = listar_projetos(url, email, administradores)
    except Exception as erro:
        _erro(erro)
        return
    if email in administradores:
        with st.expander("Criar projeto", expanded=not projetos):
            with st.form("criar_projeto"):
                nome = st.text_input("Nome do projeto", placeholder="Ex.: UTE Vale do Açu · RevK")
                descricao = st.text_area("Descrição (opcional)")
                if st.form_submit_button("Criar projeto"):
                    try:
                        criar_projeto(url, email, administradores, nome, descricao)
                        st.toast("Projeto criado.")
                        st.rerun()
                    except Exception as erro:
                        _erro(erro)
    if not projetos:
        st.info("Nenhum projeto disponível para sua conta. Peça ao coordenador para cadastrá-lo.")
        return
    projeto = st.selectbox("Escolha o projeto", projetos,
                           format_func=lambda p: p["nome"], key="projeto_acompanhamento")
    projeto_id = projeto["id"]
    try:
        papel = papel_no_projeto(url, projeto_id, email, administradores)
        membros = listar_membros(url, projeto_id, email, administradores)
        itens = listar_itens(url, projeto_id, email, administradores)
    except Exception as erro:
        _erro(erro)
        return
    coordenador = papel == "Coordenador"
    st.caption(f'{projeto["descricao"]} · Seu acesso: {papel}')
    painel, atualizar, equipe, historico = st.tabs(
        ["Painel", "Atualizar BowTie", "Equipe e cadastro", "Histórico"]
    )
    with painel:
        _painel(url, projeto_id, email, administradores, projeto_id)
    with atualizar:
        _editar(url, projeto_id, email, administradores, coordenador, membros, itens)
    with equipe:
        st.subheader("Pessoas deste projeto")
        st.dataframe(pd.DataFrame(membros, columns=["email", "nome", "papel"])
                     .rename(columns={"email": "E-mail", "nome": "Nome", "papel": "Função"}),
                     hide_index=True, use_container_width=True)
        if coordenador:
            with st.form(f"novo_membro_{projeto_id}"):
                nome_pessoa = st.text_input("Nome da pessoa")
                email_pessoa = st.text_input("E-mail da conta usada para entrar no site")
                papel_pessoa = st.selectbox("Função", ["Analista", "Coordenador"])
                if st.form_submit_button("Cadastrar ou atualizar pessoa"):
                    try:
                        salvar_membro(url, projeto_id, email, administradores,
                                     email_pessoa, nome_pessoa, papel_pessoa)
                        st.toast("Pessoa cadastrada neste projeto.")
                        st.rerun()
                    except Exception as erro:
                        _erro(erro)
            _cadastro(url, projeto_id, email, administradores)
    with historico:
        try:
            eventos = listar_historico(url, projeto_id, email, administradores)
            st.dataframe(pd.DataFrame([{
                "Quando": e["registrado_em"], "Ação": e["acao"],
                "ID do cenário": e["id_cenario"], "Por": e["autor"],
                "Antes": str(e["antes"] or ""), "Depois": str(e["depois"] or ""),
            } for e in eventos]), hide_index=True, use_container_width=True)
        except Exception as erro:
            _erro(erro)
