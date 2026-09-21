import re
from collections import Counter, defaultdict
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from comparacao import candidato_bowtie, inventariar_salvaguardas_bowtie


AZUL = "1F4E78"
AZUL_CLARO = "D9EAF7"
VERDE = "548235"
VERDE_CLARO = "E2F0D9"
VERMELHO_CLARO = "FCE4D6"
AMARELO_CLARO = "FFF2CC"
LARANJA = "F4B183"
BRANCO = "FFFFFF"
BORDA = Side(style="thin", color="B7B7B7")


APR_HEADERS = [
    "ID", "Nó", "Sistema", "Trecho de análise", "Cenário", "Perigo", "Causas",
    "Possíveis efeitos", "Modo de detecção", "Salvaguardas preventivas (SP)",
    "Salvaguardas mitigadoras (SM)", "Salvaguardas sem tipo", "F", "Sev S",
    "Sev P", "Sev M", "Sev I", "Risco S", "Risco P", "Risco M", "Risco I",
    "Linha C (com salvaguardas)", "Observações / Recomendações",
    "Candidato Bow Tie (sev IV/V)", "Página PDF",
]

ALTERACOES_HEADERS = [
    "#", "Tipo de alteração", "ID C", "ID K", "Sistema", "Perigo",
    "Rev C (antes)", "Rev K (depois)", "Detalhe / comentário",
]

SALVAGUARDAS_HEADERS = [
    "Revisão", "ID", "Nó", "Sistema", "Perigo", "Seção na APR", "Código",
    "Descrição da salvaguarda", "Nome da barreira", "Tipo Bow Tie", "Confiança",
    "Justificativa", "Alerta IT", "Candidato Bow Tie (cenário)",
    "Validação (preencher)", "Barreira final (preencher)",
]

ENQUADRAMENTO_HEADERS = [
    "ID K", "Nó", "Sistema", "Perigo (Evento Topo)", "Ameaças (causas APR)",
    "Consequências", "F", "Sev S/P/M/I", "Risco S/P/M/I", "Candidato Bow Tie",
    "B1 Contenção Primária", "B2 Controle Básico", "B3 Alarmes/Interv. Humana",
    "B4 Intertravamento", "B5 Alívio", "B6 Proteção Pós-Liberação",
    "B7 Fonte de Ignição", "B8 Resp. Emerg. Operação", "B9 Resp. Emerg. Brigada",
    "B10 Evacuação/Abandono", "Controles de degradação / não barreira", "A avaliar",
    "Lacunas / pontos de atenção para a validação",
]

MATRIZ_HEADERS = [
    "ID K", "Sistema", "Perigo", "Sev máx (S/P/M)", "Candidato Bow Tie",
    "B1", "B2", "B3", "B4", "B5", "B6", "B7", "B8", "B9", "B10",
    "Preventivas (B1–B5)", "Mitigadoras (B6–B10)",
    "Ctrl. degradação / não barreira", "A avaliar",
]

CANDIDATOS_HEADERS = [
    "Situação Bow Tie",
    "ID C",
    "ID K",
    "Nó",
    "Sistema",
    "Nome do desvio",
    "Situação do cenário",
    "Candidato na APR antiga",
    "Candidato na APR atual",
    "Severidade antiga S/P/M/I",
    "Severidade atual S/P/M/I",
    "Dimensões críticas atuais",
    "Critério",
]

STATUS_CLARO = {
    "MANTIDO / ALTERADO": "NOME DO DESVIO MANTIDO",
    "RENOMEADO / ALTERADO": "NOME DO DESVIO ALTERADO",
    "REMOVIDO": "DESVIO REMOVIDO",
    "NOVO": "DESVIO NOVO",
}


def _seguro(valor):
    if valor is None:
        return ""
    if isinstance(valor, (int, float)):
        return valor
    texto = str(valor)
    return "'" + texto if texto.startswith("=") else texto


def _sem_marcadores(texto):
    linhas = []
    for linha in str(texto or "").splitlines():
        linha = re.sub(r"^\s*[-•]\s*", "", linha).strip()
        if linha:
            linhas.append(linha)
    return "\n".join(linhas)


def _chave_id(valor):
    achado = re.search(r"N\s*(\d+)\s*[.]\s*(\d+)", str(valor or ""), re.I)
    if achado:
        return int(achado.group(1)), int(achado.group(2)), str(valor or "")
    numeros = [int(numero) for numero in re.findall(r"\d+", str(valor or ""))]
    return (
        numeros[0] if numeros else 10**9,
        numeros[1] if len(numeros) > 1 else 10**9,
        str(valor or ""),
    )


def _ordenar_cenarios(cenarios):
    return sorted(cenarios, key=lambda cenario: _chave_id(cenario.get("ID", "")))


def _codigo_revisao(documento, padrao):
    texto = " ".join(
        str(documento.get(campo, "") or "")
        for campo in ("relatorio", "nome_arquivo")
    )
    achado = re.search(
        r"(?:\.\s*REV|\bREV(?:IS[AÃ]O)?\s*)[-_. ]*([A-Z0-9]+)",
        texto,
        re.I,
    )
    if not achado:
        return padrao
    codigo = re.sub(r"[^A-Za-z0-9]", "", achado.group(1)).upper()
    return codigo or padrao


def _nome_aba_apr(codigo):
    return f"APR_Rev{codigo}"[:31]


def _cabecalho(celula, preenchimento=AZUL):
    celula.font = Font(name="Aptos", size=10, bold=True, color=BRANCO)
    celula.fill = PatternFill("solid", fgColor=preenchimento)
    celula.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    celula.border = Border(left=BORDA, right=BORDA, top=BORDA, bottom=BORDA)


def _estilo_corpo(celula):
    celula.font = Font(name="Aptos", size=9)
    celula.alignment = Alignment(vertical="top", wrap_text=True)
    celula.border = Border(bottom=Side(style="hair", color="D9D9D9"))


def _configurar_pagina(ws, congelar, filtro, orientacao="landscape"):
    ws.freeze_panes = congelar
    ws.auto_filter.ref = filtro
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = orientacao
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _escrever_tabela(ws, headers, linhas, congelar, larguras, header_row=1):
    for coluna, header in enumerate(headers, 1):
        _cabecalho(ws.cell(header_row, coluna, header))
    ws.row_dimensions[header_row].height = 42

    for numero_linha, linha in enumerate(linhas, header_row + 1):
        for numero_coluna, header in enumerate(headers, 1):
            celula = ws.cell(
                numero_linha, numero_coluna, _seguro(linha.get(header, ""))
            )
            _estilo_corpo(celula)

    ultima = max(header_row + 1, header_row + len(linhas))
    _configurar_pagina(
        ws,
        congelar,
        f"A{header_row}:{get_column_letter(len(headers))}{ultima}",
    )
    for coluna, largura in larguras.items():
        ws.column_dimensions[coluna].width = largura
    ws.print_title_rows = f"{header_row}:{header_row}"
    return ultima


def _linhas_apr(cenarios):
    linhas = []
    for cenario in _ordenar_cenarios(cenarios):
        linha = {header: cenario.get(header, "") for header in APR_HEADERS}
        linha["Linha C (com salvaguardas)"] = cenario.get(
            "Linha C (com salvaguardas)", ""
        )
        linha["Candidato Bow Tie (sev IV/V)"] = (
            "SIM" if candidato_bowtie(cenario) else "Não"
        )
        linhas.append(linha)
    return linhas


def _criar_apr(wb, nome, cenarios):
    ws = wb.create_sheet(nome)
    larguras = {
        "A": 9, "B": 6, "C": 27, "D": 45, "E": 9, "F": 35,
        "G": 45, "H": 40, "I": 37, "J": 48, "K": 48, "L": 27,
        "M": 6, "N": 7, "O": 7, "P": 7, "Q": 7, "R": 8, "S": 8,
        "T": 8, "U": 8, "V": 24, "W": 42, "X": 15, "Y": 10,
    }
    ultima = _escrever_tabela(
        ws, APR_HEADERS, _linhas_apr(cenarios), "G2", larguras
    )
    for linha in range(2, ultima + 1):
        if ws.cell(linha, 24).value == "SIM":
            ws.cell(linha, 24).fill = PatternFill("solid", fgColor=LARANJA)
            ws.cell(linha, 24).font = Font(name="Aptos", size=9, bold=True)
    return ws


def _alteracoes_por_par(alteracoes):
    grupos = defaultdict(list)
    for item in alteracoes:
        grupos[(item.get("ID antiga", ""), item.get("ID atualizada", ""))].append(item)
    return grupos


def _mudanca(valores_antigos, valores_novos):
    antigo = "/".join(str(v or "") for v in valores_antigos)
    novo = "/".join(str(v or "") for v in valores_novos)
    return f"{antigo} → {novo}" if antigo != novo else ""


def _resumir_itens(alteracoes, prefixo, lado):
    resultado = []
    for item in alteracoes:
        tipo = str(item.get("Tipo de alteração", ""))
        if not tipo.startswith(prefixo):
            continue
        secao = ""
        achado = re.search(r"\(([^)]+)\)", tipo)
        if achado:
            secao = f"[{achado.group(1)}] "
        valor = item.get(lado, "")
        if valor:
            resultado.append(f"{secao}{valor}")
    return "\n".join(resultado)


def _linhas_comparativo(antigos, atuais, comparacao):
    mapa_c = {cenario["ID"]: cenario for cenario in antigos}
    mapa_k = {cenario["ID"]: cenario for cenario in atuais}
    grupos = _alteracoes_por_par(comparacao["alteracoes"])
    linhas = []

    relacoes = sorted(
        comparacao["correspondencias"],
        key=lambda item: _chave_id(item.get("ID atualizada") or item.get("ID antiga")),
    )
    for relacao in relacoes:
        id_c = relacao.get("ID antiga", "")
        id_k = relacao.get("ID atualizada", "")
        antigo = mapa_c.get(id_c, {})
        atual = mapa_k.get(id_k, {})
        alteracoes = grupos.get((id_c, id_k), [])
        status = relacao.get("Status", "")

        causas_incluidas = [
            item.get("Depois", "")
            for item in alteracoes
            if item.get("Tipo de alteração") == "Causa incluída" and item.get("Depois")
        ]
        causas_excluidas = [
            item.get("Antes", "")
            for item in alteracoes
            if item.get("Tipo de alteração") == "Causa excluída" and item.get("Antes")
        ]
        efeitos_incluidos = [
            item.get("Depois", "")
            for item in alteracoes
            if item.get("Tipo de alteração") == "Efeito incluído" and item.get("Depois")
        ]
        efeitos_excluidos = [
            item.get("Antes", "")
            for item in alteracoes
            if item.get("Tipo de alteração") == "Efeito excluído" and item.get("Antes")
        ]

        alteradas = []
        for item in alteracoes:
            if str(item.get("Tipo de alteração", "")).startswith(
                "Salvaguarda alterada"
            ):
                secao = re.search(r"\(([^)]+)\)", item["Tipo de alteração"])
                rotulo = f"[{secao.group(1)}] " if secao else ""
                alteradas.append(
                    f"{rotulo}{item.get('Antes', '')}  ⇒  {item.get('Depois', '')}"
                )

        nota = ""
        if id_c and id_k:
            nota = (
                f"{relacao.get('Correspondência', '')}; confiança "
                f"{str(relacao.get('Confiança', '')).lower()}; "
                f"similaridade {relacao.get('Similaridade', 0):.0f}%"
            )

        linha = {
            "Status": STATUS_CLARO.get(status, status),
            "ID C": id_c,
            "ID K": id_k,
            "Nó C": antigo.get("Nó", ""),
            "Nó K": atual.get("Nó", ""),
            "Sistema": atual.get("Sistema", antigo.get("Sistema", "")),
            "Perigo (C)": _sem_marcadores(antigo.get("Perigo", "")),
            "Perigo (K)": _sem_marcadores(atual.get("Perigo", "")),
            "Causas (C)": _sem_marcadores(antigo.get("Causas", "")),
            "Causas (K)": _sem_marcadores(atual.get("Causas", "")),
            "Efeitos (C)": _sem_marcadores(antigo.get("Possíveis efeitos", "")),
            "Efeitos (K)": _sem_marcadores(atual.get("Possíveis efeitos", "")),
            "Detecção (C)": _sem_marcadores(antigo.get("Modo de detecção", "")),
            "Detecção (K)": _sem_marcadores(atual.get("Modo de detecção", "")),
            "Preventivas (C)": _sem_marcadores(
                antigo.get("Salvaguardas preventivas (SP)", "")
            ),
            "Preventivas (K)": _sem_marcadores(
                atual.get("Salvaguardas preventivas (SP)", "")
            ),
            "Mitigadoras (C)": _sem_marcadores(
                antigo.get("Salvaguardas mitigadoras (SM)", "")
            ),
            "Mitigadoras (K)": _sem_marcadores(
                atual.get("Salvaguardas mitigadoras (SM)", "")
            ),
            "F (C)": antigo.get("F", ""),
            "F (K)": atual.get("F", ""),
            "Sev C S": antigo.get("Sev S", ""),
            "Sev C P": antigo.get("Sev P", ""),
            "Sev C M": antigo.get("Sev M", ""),
            "Sev C I": antigo.get("Sev I", ""),
            "Sev K S": atual.get("Sev S", ""),
            "Sev K P": atual.get("Sev P", ""),
            "Sev K M": atual.get("Sev M", ""),
            "Sev K I": atual.get("Sev I", ""),
            "Risco C S": antigo.get("Risco S", ""),
            "Risco C P": antigo.get("Risco P", ""),
            "Risco C M": antigo.get("Risco M", ""),
            "Risco C I": antigo.get("Risco I", ""),
            "Risco K S": atual.get("Risco S", ""),
            "Risco K P": atual.get("Risco P", ""),
            "Risco K M": atual.get("Risco M", ""),
            "Risco K I": atual.get("Risco I", ""),
            "Obs/Recom. (C)": antigo.get("Observações / Recomendações", ""),
            "Obs/Recom. (K)": atual.get("Observações / Recomendações", ""),
            "Δ Freq.": (
                f"{antigo.get('F', '')} → {atual.get('F', '')}"
                if antigo and atual and antigo.get("F", "") != atual.get("F", "")
                else ""
            ),
            "Δ Severidade": (
                _mudanca(
                    [antigo.get(f"Sev {x}", "") for x in "SPMI"],
                    [atual.get(f"Sev {x}", "") for x in "SPMI"],
                )
                if antigo and atual
                else ""
            ),
            "Δ Risco": (
                _mudanca(
                    [antigo.get(f"Risco {x}", "") for x in "SPMI"],
                    [atual.get(f"Risco {x}", "") for x in "SPMI"],
                )
                if antigo and atual
                else ""
            ),
            "Δ Causas": "\n".join(
                (["Incluídas: " + "; ".join(causas_incluidas)] if causas_incluidas else [])
                + (["Excluídas: " + "; ".join(causas_excluidas)] if causas_excluidas else [])
            ),
            "Δ Efeitos": "\n".join(
                (["Incluídos: " + "; ".join(efeitos_incluidos)] if efeitos_incluidos else [])
                + (["Excluídos: " + "; ".join(efeitos_excluidos)] if efeitos_excluidos else [])
            ),
            "Salvaguardas incluídas (K)": _resumir_itens(
                alteracoes, "Salvaguarda incluída", "Depois"
            ),
            "Salvaguardas excluídas (C)": _resumir_itens(
                alteracoes, "Salvaguarda excluída", "Antes"
            ),
            "Salvaguardas alteradas (texto)": "\n".join(alteradas),
            "Bow Tie? (K)": (
                "SIM" if atual and candidato_bowtie(atual) else ("Não" if atual else "")
            ),
            "Nota de correspondência": nota,
        }
        linhas.append(linha)
    return linhas


def _criar_comparativo(wb, antigos, atuais, comparacao):
    ws = wb.create_sheet("Comparativo")
    top = [
        "Status", "ID C", "ID K", "Nó C", "Nó K", "Sistema", "Perigo (C)",
        "Perigo (K)", "Causas (C)", "Causas (K)", "Efeitos (C)", "Efeitos (K)",
        "Detecção (C)", "Detecção (K)", "Preventivas (C)", "Preventivas (K)",
        "Mitigadoras (C)", "Mitigadoras (K)", "Frequência", "", "Severidade ( C )",
        "", "", "", "Severidade (K)", "", "", "", "Risco ( C )", "", "", "",
        "Risco (K)", "", "", "", "Obs/Recom. (C)", "Obs/Recom. (K)", "Δ Freq.",
        "Δ Severidade", "Δ Risco", "Δ Causas", "Δ Efeitos",
        "Salvaguardas incluídas (K)", "Salvaguardas excluídas (C)",
        "Salvaguardas alteradas (texto)", "Bow Tie? (K)", "Nota de correspondência",
    ]
    sub = [""] * 48
    sub[18:20] = ["F (C)", "F (K)"]
    sub[20:24] = ["S", "P", "M", "I"]
    sub[24:28] = ["S", "P", "M", "I"]
    sub[28:32] = ["S", "P", "M", "I"]
    sub[32:36] = ["S", "P", "M", "I"]

    for coluna, valor in enumerate(top, 1):
        if valor:
            ws.cell(1, coluna, valor)
    for coluna, valor in enumerate(sub, 1):
        if valor:
            ws.cell(2, coluna, valor)

    for coluna in range(1, 19):
        ws.merge_cells(
            start_row=1, start_column=coluna, end_row=2, end_column=coluna
        )
    for inicio, fim in ((19, 20), (21, 24), (25, 28), (29, 32), (33, 36)):
        ws.merge_cells(start_row=1, start_column=inicio, end_row=1, end_column=fim)
    for coluna in range(37, 49):
        ws.merge_cells(
            start_row=1, start_column=coluna, end_row=2, end_column=coluna
        )

    for linha in (1, 2):
        for coluna in range(1, 49):
            preenchimento = AZUL
            if (
                coluna in (3, 5, 8, 10, 12, 14, 16, 18, 20)
                or 25 <= coluna <= 28
                or 33 <= coluna <= 36
            ):
                preenchimento = VERDE
            _cabecalho(ws.cell(linha, coluna), preenchimento)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 25

    chaves = [
        "Status", "ID C", "ID K", "Nó C", "Nó K", "Sistema", "Perigo (C)",
        "Perigo (K)", "Causas (C)", "Causas (K)", "Efeitos (C)", "Efeitos (K)",
        "Detecção (C)", "Detecção (K)", "Preventivas (C)", "Preventivas (K)",
        "Mitigadoras (C)", "Mitigadoras (K)", "F (C)", "F (K)", "Sev C S",
        "Sev C P", "Sev C M", "Sev C I", "Sev K S", "Sev K P", "Sev K M",
        "Sev K I", "Risco C S", "Risco C P", "Risco C M", "Risco C I",
        "Risco K S", "Risco K P", "Risco K M", "Risco K I", "Obs/Recom. (C)",
        "Obs/Recom. (K)", "Δ Freq.", "Δ Severidade", "Δ Risco", "Δ Causas",
        "Δ Efeitos", "Salvaguardas incluídas (K)", "Salvaguardas excluídas (C)",
        "Salvaguardas alteradas (texto)", "Bow Tie? (K)", "Nota de correspondência",
    ]
    linhas = _linhas_comparativo(antigos, atuais, comparacao)
    for numero_linha, linha in enumerate(linhas, 3):
        status = linha["Status"]
        preenchimento = (
            VERDE_CLARO
            if status == "DESVIO NOVO"
            else VERMELHO_CLARO
            if status == "DESVIO REMOVIDO"
            else AMARELO_CLARO
        )
        for numero_coluna, chave in enumerate(chaves, 1):
            celula = ws.cell(
                numero_linha, numero_coluna, _seguro(linha.get(chave, ""))
            )
            _estilo_corpo(celula)
            celula.fill = PatternFill("solid", fgColor=preenchimento)

    ultima = 2 + len(linhas)
    _configurar_pagina(ws, "R3", f"A2:AV{ultima}")
    ws.print_title_rows = "1:2"
    widths = {
        "A": 22, "B": 8, "C": 8, "D": 6, "E": 6, "F": 26,
        "G": 33, "H": 33, "I": 38, "J": 38, "K": 35, "L": 35,
        "M": 32, "N": 32, "O": 45, "P": 45, "Q": 45, "R": 45,
        "S": 7, "T": 7,
    }
    for coluna in range(21, 37):
        widths[get_column_letter(coluna)] = 6
    for coluna in range(37, 49):
        widths[get_column_letter(coluna)] = 32
    for coluna, largura in widths.items():
        ws.column_dimensions[coluna].width = largura
    return ws


def _linhas_alteracoes(alteracoes):
    linhas = []
    ordenadas = sorted(
        alteracoes,
        key=lambda item: (
            _chave_id(item.get("ID atualizada") or item.get("ID antiga")),
            item.get("Tipo de alteração", ""),
        ),
    )
    for numero, item in enumerate(ordenadas, 1):
        linhas.append(
            {
                "#": numero,
                "Tipo de alteração": item.get("Tipo de alteração", ""),
                "ID C": item.get("ID antiga", ""),
                "ID K": item.get("ID atualizada", ""),
                "Sistema": item.get("Sistema", ""),
                "Perigo": _sem_marcadores(item.get("Perigo", "")),
                "Rev C (antes)": item.get("Antes", ""),
                "Rev K (depois)": item.get("Depois", ""),
                "Detalhe / comentário": item.get("Detalhe / comentário", ""),
            }
        )
    return linhas


def _criar_alteracoes(wb, alteracoes):
    ws = wb.create_sheet("Alterações")
    larguras = {
        "A": 7, "B": 31, "C": 9, "D": 9, "E": 28,
        "F": 37, "G": 48, "H": 48, "I": 40,
    }
    _escrever_tabela(
        ws, ALTERACOES_HEADERS, _linhas_alteracoes(alteracoes), "C2", larguras
    )
    return ws


def _descricao_sem_codigo(texto):
    return re.sub(
        r"^(?:SP|SM)\s*\d+\s*[-–:]?\s*", "", str(texto or ""), flags=re.I
    ).strip()


def _inventario(cenarios, revisao):
    linhas = []
    for item in inventariar_salvaguardas_bowtie(
        _ordenar_cenarios(cenarios), revisao
    ):
        linhas.append(
            {
                "Revisão": revisao,
                "ID": item.get("ID", ""),
                "Nó": item.get("Nó", ""),
                "Sistema": item.get("Sistema", ""),
                "Perigo": _sem_marcadores(item.get("Perigo", "")),
                "Seção na APR": item.get("Seção na APR", ""),
                "Código": item.get("Código", ""),
                "Descrição da salvaguarda": _descricao_sem_codigo(
                    item.get("Descrição", "")
                ),
                "Nome da barreira": item.get("Nome da barreira", ""),
                "Tipo Bow Tie": item.get("Tipo Bow Tie", ""),
                "Confiança": item.get("Confiança", ""),
                "Justificativa": item.get("Justificativa", ""),
                "Alerta IT": item.get("Alerta", ""),
                "Candidato Bow Tie (cenário)": item.get(
                    "Candidato Bow Tie mínimo", ""
                ),
                "Validação (preencher)": "",
                "Barreira final (preencher)": "",
                "_Código Bow Tie": item.get("Código Bow Tie", ""),
                "_Descrição original": item.get("Descrição", ""),
            }
        )
    return linhas


def _criar_salvaguardas(wb, inventario):
    ws = wb.create_sheet("Salvaguardas")
    larguras = {
        "A": 9, "B": 9, "C": 6, "D": 27, "E": 34, "F": 15,
        "G": 10, "H": 50, "I": 31, "J": 18, "K": 14, "L": 58,
        "M": 46, "N": 18, "O": 22, "P": 23,
    }
    _escrever_tabela(
        ws, SALVAGUARDAS_HEADERS, inventario, "H2", larguras
    )
    return ws


def _linhas_candidatos(evolucao):
    relevantes = [
        item
        for item in evolucao
        if item.get("Situação Bow Tie") != "NÃO CANDIDATO"
    ]
    relevantes.sort(
        key=lambda item: _chave_id(
            item.get("ID atualizada") or item.get("ID antiga")
        )
    )
    linhas = []
    for item in relevantes:
        linhas.append(
            {
                "Situação Bow Tie": item.get("Situação Bow Tie", ""),
                "ID C": item.get("ID antiga", ""),
                "ID K": item.get("ID atualizada", ""),
                "Nó": item.get("Nó", ""),
                "Sistema": item.get("Sistema", ""),
                "Nome do desvio": _sem_marcadores(item.get("Perigo", "")),
                "Situação do cenário": STATUS_CLARO.get(
                    item.get("Status do cenário", ""),
                    item.get("Status do cenário", ""),
                ),
                "Candidato na APR antiga": item.get(
                    "Candidato na APR antiga", ""
                ),
                "Candidato na APR atual": item.get(
                    "Candidato na APR atualizada", ""
                ),
                "Severidade antiga S/P/M/I": item.get(
                    "Severidade antiga S/P/M/I", ""
                ),
                "Severidade atual S/P/M/I": item.get(
                    "Severidade atualizada S/P/M/I", ""
                ),
                "Dimensões críticas atuais": item.get(
                    "Dimensões críticas atuais", ""
                ),
                "Critério": item.get("Critério", ""),
            }
        )
    return linhas


def _criar_candidatos(wb, evolucao):
    ws = wb.create_sheet("Candidatos_BowTie")
    larguras = {
        "A": 25,
        "B": 9,
        "C": 9,
        "D": 7,
        "E": 30,
        "F": 39,
        "G": 29,
        "H": 20,
        "I": 20,
        "J": 22,
        "K": 22,
        "L": 30,
        "M": 55,
    }
    linhas = _linhas_candidatos(evolucao)
    ultima = _escrever_tabela(
        ws, CANDIDATOS_HEADERS, linhas, "F2", larguras
    )
    cores = {
        "NOVO CANDIDATO": VERDE_CLARO,
        "PERMANECE CANDIDATO": AMARELO_CLARO,
        "DEIXOU DE SER CANDIDATO": VERMELHO_CLARO,
        "CANDIDATO REMOVIDO": VERMELHO_CLARO,
    }
    for linha in range(2, ultima + 1):
        cor = cores.get(ws.cell(linha, 1).value)
        if cor:
            for coluna in range(1, len(CANDIDATOS_HEADERS) + 1):
                ws.cell(linha, coluna).fill = PatternFill("solid", fgColor=cor)
    return ws


def _severidade_maxima(cenario):
    ordem = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5}
    valores = [str(cenario.get(f"Sev {dim}", "") or "") for dim in "SPM"]
    return max(
        valores,
        key=lambda valor: ordem.get(valor.strip().upper(), 0),
        default="",
    )


def _inventario_por_id(inventario):
    resultado = defaultdict(list)
    for item in inventario:
        resultado[item["ID"]].append(item)
    return resultado


def _linhas_enquadramento(cenarios, inventario):
    por_id = _inventario_por_id(inventario)
    colunas = {
        "B1": "B1 Contenção Primária",
        "B2": "B2 Controle Básico",
        "B3": "B3 Alarmes/Interv. Humana",
        "B4": "B4 Intertravamento",
        "B5": "B5 Alívio",
        "B6": "B6 Proteção Pós-Liberação",
        "B7": "B7 Fonte de Ignição",
        "B8": "B8 Resp. Emerg. Operação",
        "B9": "B9 Resp. Emerg. Brigada",
        "B10": "B10 Evacuação/Abandono",
    }
    linhas = []
    for cenario in _ordenar_cenarios(cenarios):
        grupos = defaultdict(list)
        alertas = []
        for item in por_id.get(cenario["ID"], []):
            codigo = item.get("_Código Bow Tie", "AV")
            grupos[codigo].append(item.get("_Descrição original", ""))
            if item.get("Alerta IT"):
                alertas.append(item["Alerta IT"])

        candidato = candidato_bowtie(cenario)
        lacunas = list(dict.fromkeys(alertas))
        if candidato:
            lacunas.insert(
                0,
                "Cenário de acidente maior (sev IV/V): candidato a Bow Tie (Guia, item 9).",
            )
            if not grupos.get("B1"):
                lacunas.append(
                    "B1 sem elemento explícito: mapear a contenção primária do trecho."
                )
            if not grupos.get("B10"):
                lacunas.append(
                    "B10 sem elemento explícito: confirmar evacuação/abandono no PRE."
                )
        consequencias = str(cenario.get("Possíveis efeitos", "")).lower()
        if (
            ("incênd" in consequencias or "explos" in consequencias)
            and not grupos.get("B7")
        ):
            lacunas.append(
                "Consequência com incêndio/explosão: confirmar controles de fonte de ignição (B7)."
            )
        observacao = str(
            cenario.get("Observações / Recomendações", "") or ""
        ).strip()
        if observacao:
            lacunas.append(f"Obs/Recomendação da APR: {observacao}")

        linha = {
            "ID K": cenario.get("ID", ""),
            "Nó": cenario.get("Nó", ""),
            "Sistema": cenario.get("Sistema", ""),
            "Perigo (Evento Topo)": _sem_marcadores(cenario.get("Perigo", "")),
            "Ameaças (causas APR)": _sem_marcadores(cenario.get("Causas", "")),
            "Consequências": _sem_marcadores(
                cenario.get("Possíveis efeitos", "")
            ),
            "F": cenario.get("F", ""),
            "Sev S/P/M/I": "/".join(
                str(cenario.get(f"Sev {x}", "") or "") for x in "SPMI"
            ),
            "Risco S/P/M/I": "/".join(
                str(cenario.get(f"Risco {x}", "") or "") for x in "SPMI"
            ),
            "Candidato Bow Tie": "SIM" if candidato else "Não",
            "Controles de degradação / não barreira": "\n".join(
                grupos.get("CD", []) + grupos.get("NB", [])
            ),
            "A avaliar": "\n".join(grupos.get("AV", [])),
            "Lacunas / pontos de atenção para a validação": "\n".join(
                dict.fromkeys(lacunas)
            ),
        }
        for codigo, coluna in colunas.items():
            linha[coluna] = "\n".join(grupos.get(codigo, []))
        linhas.append(linha)
    return linhas


def _criar_enquadramento(wb, cenarios, inventario):
    ws = wb.create_sheet("Enquadramento_BowTie")
    larguras = {
        "A": 9, "B": 6, "C": 28, "D": 34, "E": 40, "F": 36, "G": 6,
        "H": 14, "I": 14, "J": 15, "K": 29, "L": 28, "M": 31,
        "N": 29, "O": 27, "P": 31, "Q": 26, "R": 31, "S": 31,
        "T": 29, "U": 34, "V": 32, "W": 58,
    }
    ultima = _escrever_tabela(
        ws,
        ENQUADRAMENTO_HEADERS,
        _linhas_enquadramento(cenarios, inventario),
        "E2",
        larguras,
    )
    for linha in range(2, ultima + 1):
        if ws.cell(linha, 10).value == "SIM":
            ws.cell(linha, 10).fill = PatternFill("solid", fgColor=LARANJA)
            ws.cell(linha, 10).font = Font(name="Aptos", size=9, bold=True)
    return ws


def _linhas_matriz(cenarios, inventario):
    por_id = _inventario_por_id(inventario)
    linhas = []
    for cenario in _ordenar_cenarios(cenarios):
        contagens = Counter(
            item.get("_Código Bow Tie", "AV")
            for item in por_id.get(cenario["ID"], [])
        )
        linha = {
            "ID K": cenario.get("ID", ""),
            "Sistema": cenario.get("Sistema", ""),
            "Perigo": _sem_marcadores(cenario.get("Perigo", "")),
            "Sev máx (S/P/M)": _severidade_maxima(cenario),
            "Candidato Bow Tie": "SIM" if candidato_bowtie(cenario) else "Não",
        }
        for codigo in [f"B{x}" for x in range(1, 11)]:
            linha[codigo] = contagens.get(codigo, 0) or ""
        linha["Preventivas (B1–B5)"] = sum(
            contagens.get(f"B{x}", 0) for x in range(1, 6)
        )
        linha["Mitigadoras (B6–B10)"] = sum(
            contagens.get(f"B{x}", 0) for x in range(6, 11)
        )
        linha["Ctrl. degradação / não barreira"] = (
            contagens.get("CD", 0) + contagens.get("NB", 0) or ""
        )
        linha["A avaliar"] = contagens.get("AV", 0) or ""
        linhas.append(linha)
    return linhas


def _criar_matriz(wb, cenarios, inventario):
    ws = wb.create_sheet("Matriz_Barreiras")
    larguras = {"A": 9, "B": 29, "C": 38, "D": 16, "E": 16}
    for coluna in range(6, 16):
        larguras[get_column_letter(coluna)] = 7
    larguras.update({"P": 20, "Q": 21, "R": 22, "S": 14})
    linhas = _linhas_matriz(cenarios, inventario)
    ultima = _escrever_tabela(ws, MATRIZ_HEADERS, linhas, "D2", larguras)
    total = ultima + 1
    ws.cell(total, 1, "TOTAL")
    ws.cell(total, 5, f'=COUNTIF(E2:E{ultima},"SIM")')
    for coluna in range(6, 20):
        letra = get_column_letter(coluna)
        ws.cell(total, coluna, f"=SUM({letra}2:{letra}{ultima})")
    for coluna in range(1, 20):
        _cabecalho(ws.cell(total, coluna), AZUL)
    ws.auto_filter.ref = f"A1:S{ultima}"
    for linha in range(2, ultima + 1):
        if ws.cell(linha, 5).value == "SIM":
            ws.cell(linha, 5).fill = PatternFill("solid", fgColor=LARANJA)
            ws.cell(linha, 5).font = Font(name="Aptos", size=9, bold=True)
    return ws


def _criar_leia_me(wb, antiga, atualizada, codigo_c, codigo_k):
    ws = wb.active
    ws.title = "Leia-me"
    titulo = (
        f"Validação APR – Rev {codigo_c} × Rev {codigo_k} – Enquadramento Bow Tie"
    )
    ws.cell(1, 1, titulo)
    ws.cell(1, 1).font = Font(
        name="Aptos Display", size=16, bold=True, color=AZUL
    )
    linhas = [
        (
            "Objetivo",
            "Comparar lado a lado as duas revisões da APR, listar todas as mudanças e indicar os cenários candidatos a Bow Tie.",
        ),
        (
            "Fontes",
            f"Rev {codigo_c}: {antiga.get('relatorio', '')} (aprovação {antiga.get('aprovacao', '')}) – {len(antiga['cenarios'])} cenários. Rev {codigo_k}: {atualizada.get('relatorio', '')} (aprovação {atualizada.get('aprovacao', '')}) – {len(atualizada['cenarios'])} cenários.",
        ),
        (
            "Identificador do cenário",
            "O ID e a numeração do nó/cenário são apenas referências de localização. A correspondência prioriza o nome do desvio/perigo e usa causas, efeitos, trecho e sistema como apoio; renumerações não criam falsas mudanças.",
        ),
        (
            f"Aba APR_Rev{codigo_c} / APR_Rev{codigo_k}",
            "Conversão das APRs para Excel, um cenário por linha, com frequência, severidade, risco e indicação de candidato a Bow Tie.",
        ),
        (
            "Aba Comparativo",
            "Cenários lado a lado. Sem par = REMOVIDO ou NOVO. As colunas Δ mostram mudanças de frequência, severidade, risco, causas, efeitos e salvaguardas.",
        ),
        (
            "Aba Candidatos_BowTie",
            "Lista os candidatos atuais, os novos candidatos e os cenários que deixaram de atender ao critério, sempre em ordem de nó.",
        ),
        (
            "Aba Alterações",
            "Lista analítica, uma linha por alteração detectada. Use o filtro da coluna Tipo de alteração.",
        ),
        (
            "Aba Salvaguardas",
            "Inventário de detecções e salvaguardas, com sugestão de enquadramento B1–B10, confiança, justificativa e campos para validação.",
        ),
        (
            "Aba Enquadramento_BowTie",
            "Visão por cenário da revisão vigente, com salvaguardas agrupadas por barreira e lacunas para validação.",
        ),
        (
            "Aba Matriz_Barreiras",
            "Quantidade de elementos B1–B10 por cenário. Células vazias destacam barreiras sem elemento identificado.",
        ),
        ("Aba Resumo", "Indicadores consolidados por revisão, tipo de alteração, barreira e nó."),
        (
            "Critério Bow Tie",
            "Candidato quando há severidade IV ou V em Segurança/Pessoas, Patrimônio ou Meio Ambiente, conforme o Guia Bow Tie rev.7.",
        ),
        (
            "Status do enquadramento",
            "A classificação de barreiras é uma sugestão automática. Itens com confiança média/baixa, alertas e lacunas devem ser confirmados pela equipe técnica.",
        ),
        (
            "Legenda de cores",
            "Azul = revisão antiga · Verde = revisão atual · Verde claro = NOVO · Vermelho claro = REMOVIDO · Amarelo = ALTERADO · Laranja = candidato a Bow Tie.",
        ),
    ]
    for linha, (rotulo, texto) in enumerate(linhas, 3):
        ws.cell(linha, 1, rotulo)
        ws.cell(linha, 2, texto)
        ws.cell(linha, 1).font = Font(
            name="Aptos", size=10, bold=True, color=AZUL
        )
        ws.cell(linha, 1).fill = PatternFill("solid", fgColor=AZUL_CLARO)
        for coluna in (1, 2):
            celula = ws.cell(linha, coluna)
            celula.alignment = Alignment(vertical="top", wrap_text=True)
            celula.border = Border(bottom=BORDA)
        ws.row_dimensions[linha].height = 42
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 108
    ws.sheet_view.showGridLines = False
    return ws


def _sistema_do_no(cenarios, no):
    for cenario in cenarios:
        if str(cenario.get("Nó", "")) == str(no):
            return _sem_marcadores(cenario.get("Sistema", ""))
    return "—"


def _criar_resumo(
    wb, antiga, atualizada, comparacao, inv_c, inv_k, codigo_c, codigo_k
):
    ws = wb.create_sheet("Resumo")
    ws.cell(1, 1, f"Resumo – APR Rev {codigo_c} × Rev {codigo_k}")
    ws.cell(1, 1).font = Font(
        name="Aptos Display", size=16, bold=True, color=AZUL
    )
    ws.column_dimensions["A"].width = 84
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18

    def secao(linha, valores):
        for coluna, valor in enumerate(valores, 1):
            ws.cell(linha, coluna, valor)
            _cabecalho(ws.cell(linha, coluna))

    secao(3, ["Indicador", f"Rev {codigo_c}", f"Rev {codigo_k}"])
    indicadores = [
        ("Cenários", len(antiga["cenarios"]), len(atualizada["cenarios"])),
        ("Itens de salvaguarda/detecção", len(inv_c), len(inv_k)),
        (
            "Cenários candidatos a Bow Tie (sev IV/V em S/P/M)",
            comparacao["resumo"].get("CANDIDATOS BOW TIE ANTIGA", 0),
            comparacao["resumo"].get("CANDIDATOS BOW TIE ATUALIZADA", 0),
        ),
    ]
    for linha, valores in enumerate(indicadores, 4):
        for coluna, valor in enumerate(valores, 1):
            ws.cell(linha, coluna, valor)
            _estilo_corpo(ws.cell(linha, coluna))

    tipos = [
        "Cenário removido", "Cenário novo", "Cenário renomeado",
        "Frequência alterada", "Severidade alterada", "Categoria de risco alterada",
        "Causa incluída", "Causa excluída", "Efeito incluído", "Efeito excluído",
        "Salvaguarda incluída (Detecção)", "Salvaguarda incluída (Preventiva)",
        "Salvaguarda incluída (Mitigadora)", "Salvaguarda excluída (Detecção)",
        "Salvaguarda excluída (Preventiva)", "Salvaguarda excluída (Mitigadora)",
        "Salvaguarda alterada (Detecção)", "Salvaguarda alterada (Preventiva)",
        "Salvaguarda alterada (Mitigadora)", "Observação/recomendação alterada",
    ]
    contagem = Counter(
        item.get("Tipo de alteração", "") for item in comparacao["alteracoes"]
    )
    secao(8, ["Tipo de alteração", "Qtde", ""])
    for linha, tipo in enumerate(tipos, 9):
        ws.cell(linha, 1, tipo)
        ws.cell(linha, 2, contagem.get(tipo, 0))
        _estilo_corpo(ws.cell(linha, 1))
        _estilo_corpo(ws.cell(linha, 2))
    ws.cell(29, 1, "Total de alterações listadas")
    ws.cell(29, 2, len(comparacao["alteracoes"]))
    _estilo_corpo(ws.cell(29, 1))
    _estilo_corpo(ws.cell(29, 2))

    secao(
        31,
        ["Enquadramento (itens) – barreira", f"Rev {codigo_c}", f"Rev {codigo_k}"],
    )
    nomes = {
        "B1": "Contenção Primária",
        "B2": "Controle Básico do Processo",
        "B3": "Alarmes Críticos e Intervenção Humana",
        "B4": "Intertravamento de Segurança",
        "B5": "Sistemas de Alívio",
        "B6": "Proteção Pós-Liberação",
        "B7": "Controle de Fonte de Ignição",
        "B8": "Resposta à Emergência da Operação",
        "B9": "Resposta à Emergência da Brigada",
        "B10": "Evacuação, Resgate e Abandono",
        "CD": "Controle de degradação",
        "NB": "Não se enquadra como barreira",
        "AV": "A avaliar pelo validador",
    }
    cont_c = Counter(item.get("_Código Bow Tie", "") for item in inv_c)
    cont_k = Counter(item.get("_Código Bow Tie", "") for item in inv_k)
    for linha, codigo in enumerate(
        [f"B{x}" for x in range(1, 11)] + ["CD", "NB", "AV"], 32
    ):
        ws.cell(linha, 1, f"{codigo} – {nomes[codigo]}")
        ws.cell(linha, 2, cont_c.get(codigo, 0))
        ws.cell(linha, 3, cont_k.get(codigo, 0))
        for coluna in range(1, 4):
            _estilo_corpo(ws.cell(linha, coluna))

    secao(46, ["Cenários por nó", f"Rev {codigo_c}", f"Rev {codigo_k}"])
    for no in range(1, 11):
        linha = 46 + no
        ws.cell(
            linha,
            1,
            f"Nó {no}: {codigo_c} = {_sistema_do_no(antiga['cenarios'], no)} | {codigo_k} = {_sistema_do_no(atualizada['cenarios'], no)}",
        )
        ws.cell(
            linha,
            2,
            sum(str(c.get("Nó", "")) == str(no) for c in antiga["cenarios"]),
        )
        ws.cell(
            linha,
            3,
            sum(str(c.get("Nó", "")) == str(no) for c in atualizada["cenarios"]),
        )
        for coluna in range(1, 4):
            _estilo_corpo(ws.cell(linha, coluna))
    ws.sheet_view.showGridLines = False
    return ws


def gerar_excel(antiga, atualizada, comparacao):
    codigo_c = _codigo_revisao(antiga, "C")
    codigo_k = _codigo_revisao(atualizada, "K")
    nome_c = _nome_aba_apr(codigo_c)
    nome_k = _nome_aba_apr(codigo_k)
    if nome_c == nome_k:
        nome_c, nome_k = "APR_Antiga", "APR_Atualizada"

    inv_c = _inventario(antiga["cenarios"], f"Rev {codigo_c}")
    inv_k = _inventario(atualizada["cenarios"], f"Rev {codigo_k}")

    wb = Workbook()
    _criar_leia_me(wb, antiga, atualizada, codigo_c, codigo_k)
    _criar_resumo(
        wb, antiga, atualizada, comparacao, inv_c, inv_k, codigo_c, codigo_k
    )
    _criar_candidatos(wb, comparacao["evolucao_bowtie"])
    _criar_apr(wb, nome_c, antiga["cenarios"])
    _criar_apr(wb, nome_k, atualizada["cenarios"])
    _criar_comparativo(
        wb, antiga["cenarios"], atualizada["cenarios"], comparacao
    )
    _criar_alteracoes(wb, comparacao["alteracoes"])
    _criar_salvaguardas(wb, inv_c + inv_k)
    _criar_enquadramento(wb, atualizada["cenarios"], inv_k)
    _criar_matriz(wb, atualizada["cenarios"], inv_k)

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    saida = BytesIO()
    wb.save(saida)
    saida.seek(0)
    return saida.getvalue()


def nome_arquivo_excel(antiga, atualizada):
    codigo_c = _codigo_revisao(antiga, "antiga")
    codigo_k = _codigo_revisao(atualizada, "atualizada")
    relatorio = atualizada.get("relatorio") or antiga.get("relatorio") or "APR"
    base = re.sub(r"\.REV[A-Z0-9]+", "", relatorio, flags=re.I)
    nome = (
        f"Comparacao_{base}_Rev{codigo_c}_x_Rev{codigo_k}_Validacao_BowTie.xlsx"
    )
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", nome)
