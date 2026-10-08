"""Converte um BowTie enviado em Def/Data no formato da planilha de carga.

Dados são lidos exclusivamente do arquivo recebido. O arquivo em assets fornece
apenas cabeçalhos, estilos e dimensões; seus exemplos são sempre apagados.
"""

from __future__ import annotations

import copy
import io
import re
import unicodedata
from pathlib import Path
from zipfile import BadZipFile, ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook


MODELO = Path(__file__).resolve().parent / "assets" / "modelo_planilha_carga.xlsx"
DEF = ("SITE_CODE", "AREA_CODE", "ID", "TYPE", "CODE", "DESCR")
DATA = ("SITE_CODE", "AREA_CODE", "ID", "BARRIER_DESC", "TAG", "DESCRIPTION", "FLOC", "TC_LIST")
BARREIRA = re.compile(r"B\d+\Z", re.IGNORECASE)
REFERENCIA = re.compile(r"^=?(?:'?([^'!]+)'?!)?\$?([A-Z]{1,3})\$?(\d+)$", re.IGNORECASE)


def _limpo(valor):
    if valor is None or isinstance(valor, str) and (
        not valor.strip() or valor.strip().startswith(("#N/A", "#REF!", "#VALUE!", "#DIV/0!"))
    ):
        return None
    return valor.strip() if isinstance(valor, str) else valor


def _valor(wb, valores, cel, profundidade=0):
    """Resolve links internos simples; usa cache para outras fórmulas, se houver."""
    if cel is None or profundidade > 3:
        return None
    if cel.data_type != "f":
        return _limpo(cel.value)
    formula = str(cel.value).strip()
    ref = REFERENCIA.fullmatch(formula)
    if ref:
        aba, coluna, linha = ref.groups()
        nome = aba or cel.parent.title
        if nome in wb:
            return _valor(wb, valores, wb[nome][f"{coluna.upper()}{linha}"], profundidade + 1)
    return _limpo(valores[cel.parent.title][cel.coordinate].value)


def _cabecalho(aba, linha=1):
    return {
        str(c.value).strip().upper(): c.column
        for c in aba[linha] if isinstance(c.value, str) and c.value.strip()
    }


def _campo(wb, valores, aba, coluna, linha):
    return _valor(wb, valores, aba.cell(linha, coluna)) if coluna else None


def _id(cel):
    if cel is None:
        return None
    if isinstance(cel, float) and cel.is_integer():
        return int(cel)
    texto = str(cel).strip()
    return int(texto) if texto.isdigit() else cel


def _definicoes(wb, valores, bt, site, area, ident):
    cab = _cabecalho(bt, 4)
    def fonte(tipo):
        if not all(cab.get(c) for c in ("TYPE", "DESCR")):
            return None
        for linha in range(5, bt.max_row + 1):
            valor = _campo(wb, valores, bt, cab["TYPE"], linha)
            if valor and str(valor).strip().upper() == tipo:
                return _campo(wb, valores, bt, cab["DESCR"], linha)
        return None

    risco = _valor(wb, valores, bt["C4"]) or fonte("HAZARD:")
    evento = _valor(wb, valores, bt["I7"]) or fonte("TOP EVENT:")
    saida = []
    if risco:
        saida.append([site, area, ident, "HAZARD:", None, risco])
    if evento:
        saida.append([site, area, ident, "TOP EVENT:", None, evento])

    # A representação visual do BT contém os códigos reais em B e P.
    for numero in range(7, bt.max_row + 1):
        codigo = _valor(wb, valores, bt[f"B{numero}"])
        descricao = _valor(wb, valores, bt[f"C{numero}"])
        if codigo or descricao:
            saida.append([site, area, ident, "THREAT:", codigo, descricao])
    for numero in range(7, bt.max_row + 1):
        codigo = _valor(wb, valores, bt[f"P{numero}"])
        descricao = _valor(wb, valores, bt[f"O{numero}"])
        if codigo or descricao:
            saida.append([site, area, ident, "CONSEQUENCE:", codigo, descricao])
    return saida


def _barreiras(wb, valores, site, area, ident):
    saida = []
    for aba in wb:
        if not BARREIRA.fullmatch(aba.title):
            continue
        col = _cabecalho(aba)
        if not all(col.get(chave) for chave in ("BARRIER_DESC", "TAG", "DESCRIPTION", "FLOC", "TC_LIST")):
            continue
        for numero in range(2, aba.max_row + 1):
            itens = {
                chave: _campo(wb, valores, aba, col.get(chave), numero)
                for chave in ("SITE_CODE", "AREA_CODE", "ID", "BARRIER_DESC", "TAG",
                              "DESCRIPTION", "FLOC", "TC_LIST")
            }
            # Linhas de modelo vazias ou contendo somente o código da barreira
            # não viram registros fictícios na planilha de carga.
            if not any(itens[chave] for chave in ("TAG", "DESCRIPTION", "FLOC", "TC_LIST")):
                continue
            saida.append([
                itens["SITE_CODE"] or site, itens["AREA_CODE"] or area,
                _id(itens["ID"]) if itens["ID"] is not None else ident,
                itens["BARRIER_DESC"] or aba.title,
                itens["TAG"], itens["DESCRIPTION"], itens["FLOC"], itens["TC_LIST"],
            ])
    return saida


def _avisos_consistencia(definicoes, barreiras, site, area, ident):
    avisos = []
    for campo, valor in (("SITE_CODE", site), ("AREA_CODE", area), ("ID", ident)):
        if valor is None:
            avisos.append(f"{campo} não consta do BowTie: a coluna ficou vazia onde não há valor de origem.")
    if not barreiras:
        avisos.append("Nenhuma barreira preenchida foi encontrada nas abas B1, B2 etc.; Data terá apenas o cabeçalho.")
    for campo, indice, global_val in (("SITE_CODE", 0, site), ("AREA_CODE", 1, area), ("ID", 2, ident)):
        if global_val is not None and any(
            linha[indice] is not None and str(linha[indice]) != str(global_val)
            for linha in barreiras
        ):
            avisos.append(f"Há valores diferentes de {campo} nas abas de barreiras. Confira a planilha gerada.")
    if not any(linha[3] == "HAZARD:" for linha in definicoes):
        avisos.append("Não foi encontrada a descrição do nó/perigo (HAZARD:) na aba BT.")
    if not any(linha[3] == "TOP EVENT:" for linha in definicoes):
        avisos.append("Não foi encontrado o evento topo (TOP EVENT:) na aba BT.")
    return avisos


def _preencher(aba, linhas, colunas):
    if tuple(aba.cell(1, i).value for i in range(1, len(colunas) + 1)) != colunas:
        raise ValueError(f"O modelo de carga da aba {aba.title} tem cabeçalhos inesperados.")
    estilos = [copy.copy(aba.cell(2, i)._style) for i in range(1, len(colunas) + 1)]
    altura = aba.row_dimensions[2].height
    if aba.max_row > 1:
        aba.delete_rows(2, aba.max_row - 1)
    for numero in list(aba.row_dimensions):
        if isinstance(numero, int) and numero >= 2:
            del aba.row_dimensions[numero]
    for numero, linha in enumerate(linhas, 2):
        for coluna, valor in enumerate(linha, 1):
            cel = aba.cell(numero, coluna, valor)
            cel._style = copy.copy(estilos[coluna - 1])
        if altura:
            aba.row_dimensions[numero].height = altura
    fim = max(1, len(linhas) + 1)
    aba.auto_filter.ref = f"A1:{'F' if len(colunas) == 6 else 'H'}{fim}"
    aba.print_area = aba.auto_filter.ref


def gerar_planilha_carga(conteudo: bytes):
    """Retorna (xlsx, resumo, avisos), sem alterar o BowTie original."""
    origem = io.BytesIO(conteudo)
    try:
        wb = load_workbook(origem, read_only=False, data_only=False, keep_links=False)
        valores = load_workbook(io.BytesIO(conteudo), read_only=False,
                               data_only=True, keep_links=False)
    except Exception as exc:
        raise ValueError("O arquivo enviado não é um BowTie Excel válido (.xlsx ou .xlsm).") from exc
    if "BT" not in wb:
        raise ValueError("O arquivo precisa conter a aba BT de um BowTie.")
    bt = wb["BT"]
    cab = _cabecalho(bt, 4)
    site = _campo(wb, valores, bt, cab.get("SITE_CODE"), 5)
    area = _campo(wb, valores, bt, cab.get("AREA_CODE"), 5)
    ident = _id(_campo(wb, valores, bt, cab.get("ID"), 5))
    for aba in wb:
        if not BARREIRA.fullmatch(aba.title):
            continue
        campos = _cabecalho(aba)
        for numero in range(2, aba.max_row + 1):
            if site is None:
                site = _campo(wb, valores, aba, campos.get("SITE_CODE"), numero)
            if area is None:
                area = _campo(wb, valores, aba, campos.get("AREA_CODE"), numero)
            if ident is None:
                ident = _id(_campo(wb, valores, aba, campos.get("ID"), numero))
            if site is not None and area is not None and ident is not None:
                break
    definicoes = _definicoes(wb, valores, bt, site, area, ident)
    if not definicoes:
        raise ValueError("A aba BT não contém evento topo, ameaças, consequências ou descrição do nó.")
    barreiras = _barreiras(wb, valores, site, area, ident)
    avisos = _avisos_consistencia(definicoes, barreiras, site, area, ident)
    saida = load_workbook(MODELO)
    if saida.sheetnames != ["Def", "Data"]:
        raise ValueError("O modelo de carga precisa ter apenas as abas Def e Data.")
    _preencher(saida["Def"], definicoes, DEF)
    _preencher(saida["Data"], barreiras, DATA)
    memoria = io.BytesIO()
    saida.save(memoria)
    return memoria.getvalue(), {"def": len(definicoes), "data": len(barreiras), "id": ident}, avisos


def _nome_saida(nome_origem, usados):
    base = unicodedata.normalize("NFKD", Path(nome_origem).stem)
    base = base.encode("ascii", "ignore").decode("ascii")
    base = re.sub(r"[^A-Za-z0-9_-]+", "_", base).strip("_")[:85] or "BowTie"
    candidato = f"Planilha_de_carga_{base}.xlsx"
    posicao = 2
    while candidato.casefold() in usados:
        candidato = f"Planilha_de_carga_{base}_{posicao}.xlsx"
        posicao += 1
    usados.add(candidato.casefold())
    return candidato


def gerar_planilhas_carga_zip(conteudo: bytes):
    """Converte todos os BowTies do ZIP; recusa entrega parcial em caso de erro."""
    try:
        entrada = ZipFile(io.BytesIO(conteudo))
    except (BadZipFile, OSError) as exc:
        raise ValueError("O arquivo enviado não é um ZIP válido.") from exc
    with entrada:
        arquivos = [
            info for info in entrada.infolist()
            if not info.is_dir() and not any(
                parte.startswith(".") or parte.startswith("~$")
                for parte in Path(info.filename).parts
            ) and Path(info.filename).suffix.lower() in {".xlsx", ".xlsm"}
        ]
        if not arquivos:
            raise ValueError("O ZIP não contém arquivos BowTie .xlsx ou .xlsm.")
        if len(arquivos) > 300 or sum(info.file_size for info in arquivos) > 250 * 1024 * 1024:
            raise ValueError("O lote excede 300 BowTies ou 250 MB descompactados. Divida o ZIP em lotes menores.")

        resultados, problemas, usados = [], [], set()
        memoria = io.BytesIO()
        with ZipFile(memoria, "w", compression=ZIP_DEFLATED) as saida:
            for info in arquivos:
                try:
                    planilha, resumo, avisos = gerar_planilha_carga(entrada.read(info))
                except (ValueError, RuntimeError, BadZipFile, OSError) as exc:
                    problemas.append(f"{info.filename}: {exc}")
                    continue
                nome = _nome_saida(info.filename, usados)
                saida.writestr(nome, planilha)
                resultados.append({"origem": info.filename, "arquivo": nome,
                                   "def": resumo["def"], "data": resumo["data"], "avisos": avisos})
        if problemas:
            detalhes = "\n".join(problemas[:20])
            if len(problemas) > 20:
                detalhes += f"\n... e mais {len(problemas) - 20} arquivo(s)."
            raise ValueError("Nenhum ZIP de saída foi liberado porque alguns BowTies precisam de correção:\n" + detalhes)
        return memoria.getvalue(), resultados
